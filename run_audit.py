"""
HVAC Lead Audit Engine — Main Orchestrator
=============================================
Runs the full "Find the Bleeding Neck" pipeline:
  1. Find HVAC businesses via Google Places API (or demo data / URL list)
  2. Audit each website for 13 incompetence signals
  3. Score and rank targets (0-130 scale)
  4. Generate personalized diagnostic cold emails
  5. Output everything to a formatted XLSX ready for outreach

Usage:
  python run_audit.py              # Full run (needs GOOGLE_PLACES_API_KEY env var)
  python run_audit.py --demo       # Demo mode with sample data (no API key needed)
  python run_audit.py --urls       # Audit URLs from urls.txt file
  python run_audit.py --niche "pest control"  # Override niche keyword
"""

import sys
import json
import argparse
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import GOOGLE_PLACES_API_KEY
from audit_engine import audit_batch
from email_generator import generate_batch_emails


# ── Demo data for testing without API key ──
DEMO_BUSINESSES = [
    {"name": "Acme HVAC Services", "phone": "(212) 555-0101", "website": "https://www.york.com", "address": "123 Main St, New York, NY", "city": "New York", "state": "NY", "rating": 4.2, "review_count": 87, "maps_url": ""},
    {"name": "CoolBreeze AC & Heat", "phone": "(310) 555-0202", "website": "https://www.lennox.com", "address": "456 Sunset Blvd, Los Angeles, CA", "city": "Los Angeles", "state": "CA", "rating": 3.8, "review_count": 42, "maps_url": ""},
    {"name": "Windy City Comfort", "phone": "(312) 555-0303", "website": "https://www.carrier.com", "address": "789 Michigan Ave, Chicago, IL", "city": "Chicago", "state": "IL", "rating": 4.5, "review_count": 156, "maps_url": ""},
    {"name": "Houston Air Pros", "phone": "(713) 555-0404", "website": "https://www.trane.com", "address": "321 Energy Dr, Houston, TX", "city": "Houston", "state": "TX", "rating": 3.2, "review_count": 23, "maps_url": ""},
    {"name": "Desert Cool HVAC", "phone": "(602) 555-0505", "website": "https://www.goodmanmfg.com", "address": "555 Cactus Rd, Phoenix, AZ", "city": "Phoenix", "state": "AZ", "rating": 4.0, "review_count": 65, "maps_url": ""},
]


def load_urls_from_file(filepath="urls.txt"):
    """Load business URLs from a text file. Format: name|url|city|state (one per line)."""
    businesses = []
    with open(filepath, "r") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = line.split("|")
            businesses.append({
                "name": parts[0].strip() if len(parts) > 0 else "Unknown",
                "website": parts[1].strip() if len(parts) > 1 else "",
                "city": parts[2].strip() if len(parts) > 2 else "",
                "state": parts[3].strip() if len(parts) > 3 else "",
                "phone": "", "address": "", "rating": 0, "review_count": 0, "maps_url": "",
            })
    return businesses


def build_xlsx(audit_results, emails, output_path):
    """Build a formatted XLSX with two sheets: Audit Results + Emails."""
    wb = Workbook()

    # ── Sheet 1: Audit Results ──
    ws1 = wb.active
    ws1.title = "Audit Results"

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1A1A2E")
    hot_fill = PatternFill("solid", fgColor="FFE0E0")
    warm_fill = PatternFill("solid", fgColor="FFF3CD")
    border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    headers = ["Rank", "Score", "Grade", "Business Name", "City", "State",
               "Phone", "Website", "Rating", "Reviews", "Issues Found",
               "Top Issue 1", "Top Issue 2", "Top Issue 3", "Load Time (s)"]
    ws1.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    priority_order = [
        "no_phone_visible", "no_click_to_call", "no_ssl", "slow_load",
        "no_viewport_meta", "homepage_as_landing", "no_contact_form",
        "no_analytics", "no_reviews_on_site", "no_schema_markup",
        "no_meta_description", "no_h1", "no_gtm",
    ]

    for rank, result in enumerate(audit_results, 1):
        score = result.get("score", 0)
        grade = "HOT 🔥" if score >= 80 else ("WARM 🟡" if score >= 50 else "LOW ⚪")
        issues = result.get("issues", [])
        top = [i for i in priority_order if i in issues][:3]

        row = [
            rank, score, grade,
            result.get("name", ""), result.get("city", ""), result.get("state", ""),
            result.get("phone", ""), result.get("audited_url", result.get("website", "")),
            result.get("rating", ""), result.get("review_count", ""),
            result.get("issue_count", 0),
            top[0].replace("_", " ") if len(top) > 0 else "",
            top[1].replace("_", " ") if len(top) > 1 else "",
            top[2].replace("_", " ") if len(top) > 2 else "",
            result.get("load_time", "N/A"),
        ]
        ws1.append(row)

        row_idx = rank + 1
        if score >= 80:
            for col_idx in range(1, len(headers) + 1):
                ws1.cell(row=row_idx, column=col_idx).fill = hot_fill
        elif score >= 50:
            for col_idx in range(1, len(headers) + 1):
                ws1.cell(row=row_idx, column=col_idx).fill = warm_fill

        for col_idx in range(1, len(headers) + 1):
            ws1.cell(row=row_idx, column=col_idx).border = border
            ws1.cell(row=row_idx, column=col_idx).font = Font(name="Arial", size=10)

    # Column widths
    widths = [6, 8, 10, 30, 15, 8, 18, 40, 8, 10, 14, 22, 22, 22, 12]
    for i, w in enumerate(widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # Freeze header row
    ws1.freeze_panes = "A2"

    # ── Sheet 2: Email Drafts ──
    ws2 = wb.create_sheet("Email Drafts")
    email_headers = ["Business Name", "City", "Score", "Subject Line", "Email Body", "Website", "Phone"]
    ws2.append(email_headers)

    for col_idx in range(1, len(email_headers) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for email in emails:
        ws2.append([
            email.get("business_name", ""),
            email.get("city", ""),
            email.get("score", ""),
            email.get("subject", ""),
            email.get("body", ""),
            email.get("website", ""),
            email.get("business_phone", ""),
        ])

    email_widths = [30, 15, 8, 50, 80, 40, 18]
    for i, w in enumerate(email_widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws2.freeze_panes = "A2"

    wb.save(output_path)
    print(f"\n  📊 Output saved: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="HVAC Lead Audit Engine")
    parser.add_argument("--demo", action="store_true", help="Run with demo data (no API key needed)")
    parser.add_argument("--urls", action="store_true", help="Audit URLs from urls.txt")
    parser.add_argument("--niche", type=str, help="Override niche keyword (e.g., 'pest control')")
    args = parser.parse_args()

    print("\n" + "=" * 60)
    print("  HVAC LEAD AUDIT ENGINE")
    print("  'Find the Bleeding Neck' GTM Pipeline")
    print("=" * 60)

    # Step 1: Get businesses
    if args.demo:
        print("\n  [DEMO MODE] Using sample business data\n")
        businesses = DEMO_BUSINESSES
    elif args.urls:
        print("\n  [URL MODE] Loading from urls.txt\n")
        businesses = load_urls_from_file()
    else:
        if not GOOGLE_PLACES_API_KEY:
            print("\n  ❌ No API key found.")
            print("  Set env var: export GOOGLE_PLACES_API_KEY='your-key'")
            print("  Or run: python run_audit.py --demo")
            sys.exit(1)
        from find_businesses import find_all_businesses
        keyword = args.niche if args.niche else None
        businesses = find_all_businesses(keyword=keyword)

    if not businesses:
        print("No businesses to audit. Exiting.")
        sys.exit(1)

    # Step 2: Audit websites
    print(f"\n{'=' * 60}")
    print(f"  AUDITING {len(businesses)} WEBSITES")
    print(f"{'=' * 60}\n")
    audit_results = audit_batch(businesses, delay=1.0)

    # Step 3: Generate emails
    print(f"\n{'=' * 60}")
    print(f"  GENERATING DIAGNOSTIC EMAILS")
    print(f"{'=' * 60}\n")
    emails = generate_batch_emails(audit_results, min_score=30)
    print(f"  Generated {len(emails)} emails for qualified leads")

    # Step 4: Output to XLSX
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    niche_tag = args.niche.replace(" ", "_") if args.niche else "hvac"
    output_path = f"{niche_tag}_audit_{timestamp}.xlsx"
    build_xlsx(audit_results, emails, output_path)

    # Step 5: Summary
    hot = [r for r in audit_results if r.get("score", 0) >= 80]
    warm = [r for r in audit_results if 50 <= r.get("score", 0) < 80]
    print(f"\n{'=' * 60}")
    print(f"  RESULTS SUMMARY")
    print(f"{'=' * 60}")
    print(f"  🔥 Hot leads (score ≥ 80):  {len(hot)}")
    print(f"  🟡 Warm leads (50-79):      {len(warm)}")
    print(f"  Total audited:              {len(audit_results)}")
    print(f"  Emails generated:           {len(emails)}")
    print(f"  Output file:                {output_path}")
    print(f"{'=' * 60}\n")

    # Save raw JSON for further processing
    json_path = f"audit_raw_{timestamp}.json"
    with open(json_path, "w") as f:
        json.dump(audit_results, f, indent=2, default=str)
    print(f"  Raw data: {json_path}")

    return output_path


if __name__ == "__main__":
    main()
